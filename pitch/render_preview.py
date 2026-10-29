from pathlib import Path
from io import BytesIO
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

root=Path(__file__).parent
out=root/'slide-previews'
out.mkdir(exist_ok=True)
prs=Presentation(root/'SchooKeep_Pitch_Deck.pptx')
scale=1600/prs.slide_width
pages=[]
issues=[]
def color(c,default='#142822'):
    try:return '#'+str(c.rgb)
    except:return default
for number,slide in enumerate(prs.slides,1):
    im=Image.new('RGB',(1600,900),'white');d=ImageDraw.Draw(im)
    for shape in slide.shapes:
        x,y,w,h=[round(v*scale) for v in (shape.left,shape.top,shape.width,shape.height)]
        if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
            pic=Image.open(BytesIO(shape.image.blob)).convert('RGB').resize((w,h),Image.Resampling.LANCZOS)
            im.paste(pic,(x,y));continue
        try:
            if shape.fill.type is not None:d.rectangle((x,y,x+w,y+h),fill=color(shape.fill.fore_color,'white'))
        except:pass
        if not shape.has_text_frame or not shape.text:continue
        ty=y
        for p in shape.text_frame.paragraphs:
            size=round((p.font.size.pt if p.font.size else 20)*1600/960)
            font=ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if p.font.bold else 'arial.ttf'),size)
            line='';lines=[]
            for word in p.text.split():
                trial=(line+' '+word).strip()
                if d.textlength(trial,font=font)>w and line:lines.append(line);line=word
                else:line=trial
            if line:lines.append(line)
            for line in lines:
                d.text((x,ty),line,fill=color(p.font.color),font=font)
                ty+=size*1.12
            ty+=(p.space_after.pt if p.space_after else 0)*1600/960
        if ty-y>h+22:issues.append({'slide':number,'text':shape.text[:60],'overflow_px':round(ty-y-h)})
    im.save(out/f'Slide{number:02}.png');pages.append(im)
pages[0].save(root/'SchooKeep_Pitch_Deck.pdf',save_all=True,append_images=pages[1:],resolution=120)
contact=Image.new('RGB',(1600,5*240),'#edf3f0')
for i,p in enumerate(pages):
    thumb=p.resize((384,216),Image.Resampling.LANCZOS)
    contact.paste(thumb,((i%4)*400,(i//4)*240))
contact.save(root/'Deck_Overview.jpg',quality=92)
print('Rendered',len(pages),'slides. Geometry checks:',issues)
