from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont
import json, html, zipfile

ROOT = Path(__file__).parent
SHOTS = ROOT / 'screenshots'
OUT = ROOT / 'SchooKeep_Pitch_Deck.pptx'
rows=json.loads((ROOT/'screenshots.json').read_text())
capture_count=len({r['file'] for r in rows if 'file' in r})
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
INK, GRAY, GREEN, CORAL, WHITE = '142822', '5B7069', '087F69', 'E56752', 'FFFFFF'

def box(s, x,y,w,h,color):
    q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    q.fill.solid(); q.fill.fore_color.rgb=RGBColor.from_string(color); q.line.fill.background()
    return q

def text(s, content, x,y,w,h,size=20,color=INK,bold=False):
    tf=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap=True
    tf.margin_left=tf.margin_right=0
    tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(10)
    return tf

def base(kicker,title,sub=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    box(s,0,0,.14,7.5,GREEN)
    text(s,kicker.upper(),.55,.33,11,.3,11,GREEN,True)
    text(s,title,.55,.9,12.2,1.0,32,INK,True)
    if sub:text(s,sub,.55,1.83,12.1,.7,15,GRAY)
    box(s,.55,7.02,12.2,.012,'DCE8E3')
    text(s,'SchooKeep  |  Product pitch  |  05 October 2026',.55,7.15,10,.2,9,GRAY)
    text(s,f'{len(prs.slides):02}',12.1,7.12,.6,.25,11,GREEN,True)
    return s

def phone(s,name,x,y=2.5,h=4.05,label=''):
    f=SHOTS/(name+'.png')
    if not f.exists():return
    w=h*393/852
    box(s,x-.045,y-.045,w+.09,h+.09,INK)
    s.shapes.add_picture(str(f),Inches(x),Inches(y),height=Inches(h))
    if label:text(s,label,x-.1,y+h+.13,w+.2,.42,11,GRAY,True)

def product(kicker,title,sub,points,images):
    s=base(kicker,title,sub)
    for i,(heading,body) in enumerate(points):
        yy=2.65+i*1.15
        box(s,.58,yy,.09,.55,GREEN if i%2==0 else CORAL)
        text(s,heading,.86,yy,4.5,.35,20,INK,True)
        text(s,body,.86,yy+.43,4.45,.65,15,GRAY)
    start=6.0 if len(images)<3 else 5.65
    for i,(name,label) in enumerate(images):phone(s,name,start+i*2.38,label=label)
    s.notes_slide.notes_text_frame.text='Screens captured from the Flutter release build using a local seeded Laravel backend and supplied admin account. Demo data; no claim of production use. '+sub

def columns(kicker,title,sub,items):
    s=base(kicker,title,sub)
    n=len(items); w=11.9/n
    for i,(heading,body) in enumerate(items):
        x=.65+i*w
        text(s,f'{i+1:02}',x,2.7,w-.4,.7,42,GREEN,True)
        text(s,heading,x,3.65,w-.4,.8,24,INK,True)
        text(s,body,x,4.65,w-.5,1.9,19,GRAY)
    return s

s=base('School health, connected','SchooKeep','A shared operating layer for school health, family communication, and student safety.')
text(s,'Better handoffs.\nClearer accountability.\nOne school community.',.65,2.95,5.25,2.3,30,INK,True)
text(s,'Flutter product demonstration\nLocal demo data, captured from the running app',.65,5.7,5.0,.8,15,GRAY)
phone(s,'principal-home',6.3,y=2.45,h=4.25,label='School leadership')
phone(s,'parent-app-home',9.0,y=2.45,h=4.25,label='Parent experience')

columns('The problem','Care breaks down between handoffs','The product addresses coordination gaps across the school day.',[
 ('Fragmented context','Clinic notes, consent, medication status, and pickup instructions sit in different workflows.'),
 ('Time-sensitive decisions','A health event often requires the nurse, parent, teacher, and leadership to act on the same information.'),
 ('Limited visibility','Schools need a clear operational view and a traceable record of who acted and when.')])

columns('The solution','A connected school health workspace','Role-specific mobile screens bring each team into the same operating model.',[
 ('Coordinate care','Clinic visits, dose workflows, physician review, and student health records.'),
 ('Connect families','Consent, documents, health history, medication reporting, and authorized pickups.'),
 ('Close the loop','Leadership analytics, staff permissions, school advisories, and audit views.')])

s=base('Product breadth','Built around the whole school','Eleven role experiences, with shared authentication and system states.')
roles=['Principal','Vice principal','Nurse','Physician','Parent','Teacher','Counselor','Secretary','Cafeteria','Security','Bus driver']
for i,r in enumerate(roles):
    x=.7+(i%4)*3.1; y=2.65+(i//4)*1.23
    box(s,x,y,.38,.38,GREEN if i%2==0 else CORAL)
    text(s,r,x+.55,y-.02,2.4,.6,21,INK,True)
text(s,f'{capture_count} route screenshots included in the product library.',.7,6.48,11,.4,16,GRAY)

product('Leadership','See the school health picture','A daily overview brings care activity, documents, alerts, and staff visibility together.',[
 ('An operational home','Bring today\'s priorities into one leadership screen.'),
 ('Staff and access oversight','Manage the people who deliver each part of the service.'),
 ('Analytics and reporting','Review trends and support school-level planning.')], [('principal-home','Daily overview'),('principal-analytics','Health analytics'),('principal-staff','Staff management')])

product('Clinical care','Give the nurse a focused workspace','Clinical operations span student lookup, visits, documentation, and emergency escalation.',[
 ('Student health context','Access health information alongside the student record.'),
 ('Visit documentation','Organize routine clinic activity and visit details.'),
 ('Escalation paths','Make urgent next steps visible in the workflow.')], [('nurse-clinic','Clinic visits'),('nurse-clinic-new-visit','New visit'),('nurse-clinic-emergency-escalation','Emergency escalation')])

product('Medication workflow','Make each dose a deliberate action','Daily doses, conflict states, and inventory are represented in the clinical experience.',[
 ('Plan the day','Review medication work in a dedicated dose view.'),
 ('Confirm and flag','Use confirmation and dose-conflict screens to structure the handoff.'),
 ('Track supplies','Give medication stock its own operational view.')], [('nurse-medications','Medication list'),('nurse-medications-dose-conflict','Conflict alert'),('nurse-medications-inventory','Pharmacy inventory')])

product('Physician oversight','Extend care beyond the clinic desk','A physician workspace supports protocol review, escalations, and scheduling.',[
 ('Protocol review','Provide a dedicated review area for medication protocols.'),
 ('Clinical escalation inbox','Surface cases that need additional clinical attention.'),
 ('Availability planning','Represent physician schedules in the app.')], [('physician-dashboard','Physician home'),('physician-protocols','Protocol review'),('physician-escalations','Escalation inbox')])

product('Family experience','Keep parents part of the care loop','The parent app brings health, medication, and document workflows into a familiar mobile experience.',[
 ('A family-facing home','Give guardians a clear starting point for the school day.'),
 ('Health and medication history','Keep the child\'s care activity visible to the family.'),
 ('Documents in one place','Provide an accessible path for school health paperwork.')], [('parent-app-home','Parent home'),('parent-app-health','Health history'),('parent-app-docs','Documents')])

product('Consent and handoffs','Connect a request to a response','Consent and authorized pickup screens show the link between school operations and guardians.',[
 ('Emergency response','Provide a dedicated parent consent screen.'),
 ('Authorized persons','Organize the people approved for pickup.'),
 ('QR presentation','Represent a scannable pickup identity in the family app.')], [('parent-app-emergency-consent','Consent response'),('parent-app-authorized-persons','Authorized persons'),('parent-app-full-qrcode-person-1','Pickup QR code')])

product('Classroom awareness','Bring health context into the school day','Teachers have focused views for attendance, health considerations, referrals, and exemptions.',[
 ('Attendance context','A classroom view supports daily operational awareness.'),
 ('Health considerations','Keep relevant student needs visible to the teacher.'),
 ('Referral workflow','Offer a structured route into the school clinic.')], [('teacher-attendance','Attendance'),('teacher-health-considerations','Health considerations'),('teacher-clinic-referral','Clinic referral')])

product('Student safety','Support arrival and departure handoffs','Security and transport have dedicated operational views.',[
 ('Pickup verification','Queue and manual verification screens organize gate activity.'),
 ('Transport operations','Route, boarding, and history workflows support the driver.'),
 ('Shared responsibility','Connect the school, transport staff, and authorized guardians.')], [('security-pickups','Pickup queue'),('security-manual-verification','Manual verification'),('bus-route','Bus route')])

product('Food safety','Make cafeteria decisions more informed','Allergen alerts and point-of-sale workflows extend the product beyond the clinic.',[
 ('Allergen awareness','Provide the cafeteria with dedicated alert screens.'),
 ('Daily operations','Bring point of sale into the staff experience.'),
 ('Family budget visibility','Represent cafeteria budgets in the parent app.')], [('cafeteria-alerts','Cafeteria alerts'),('cafeteria-pos','Point of sale'),('parent-app-cafeteria-budget','Family budget')])

product('School support','Support the people behind each handoff','Counselor and secretary screens broaden the operational footprint.',[
 ('Student support','Counselor records and report workflows represent ongoing care.'),
 ('Office coordination','The secretary has student, messaging, and import views.'),
 ('Triage and follow-up','Queue screens make outstanding work visible.')], [('counselor-home','Counselor dashboard'),('secretary-messages','Office messaging'),('secretary-chatbot','Chatbot queue')])

product('Governance','Make access and readiness visible','Permission, audit, and readiness views support accountable school operations.',[
 ('Role permissions','Show the access model in a dedicated matrix.'),
 ('Audit visibility','Provide a view of recorded system activity.'),
 ('Clinic readiness','Give school operations a readiness and equipment workflow.')], [('principal-permission-matrix','Permissions'),('principal-audit','Audit log'),('vice-principal-clinic-readiness','Clinic readiness')])

product('Regional product design','Designed for the school context','The app includes English/Arabic localization and contextual school system states.',[
 ('Bilingual foundation','English and Arabic strings and right-to-left layouts are present in the codebase.'),
 ('School-wide advisories','Weather advisory screens represent shared restrictions.'),
 ('Contextual operating modes','Ramadan and after-hours states are part of the experience.')], [('system-weather-advisory','Weather advisory'),('system-ramadan','Ramadan mode'),('system-after-hours','After-hours state')])

columns('Commercial hypothesis','Start with a school-level subscription','A proposed model for validation, rather than committed pricing.',[
 ('School license','An annual subscription priced by school size or student band. Validate willingness to pay in pilots.'),
 ('Implementation','Offer setup, training, and rollout support. Scope integration work separately.'),
 ('Group expansion','Expand from one school to a school group after demonstrating adoption and operating value.')])

columns('Pilot proposal','Prove the workflow before scaling','An illustrative eight-week pilot with agreed measures and owners.',[
 ('Weeks 1-2: Prepare','Map current workflows, confirm roles, review data handling, and establish a baseline.'),
 ('Weeks 3-6: Operate','Pilot clinic, consent, and parent communications with a defined cohort and support cadence.'),
 ('Weeks 7-8: Evaluate','Measure adoption, documentation completeness, consent response time, and handoff follow-through.')])

s=base('The next step','Find the right pilot school','Validate the value with the people who will use the system every day.')
text(s,'Seeking a design partner',.7,2.85,7,.65,30,GREEN,True)
text(s,'A school leadership team, nurse, and parent cohort willing to shape the rollout and define success.',.7,3.8,7.1,1.2,24,INK)
text(s,'Agree the pilot scope.\nName an operational owner.\nChoose the measures that matter.',.7,5.2,7.2,1.35,22,GRAY)
phone(s,'parent-app-home',9.5,y=2.4,h=4.3)

s=base('Evidence and scope','What this presentation demonstrates','Product screenshots are evidence of the current app experience, not commercial or clinical outcomes.')
text(s,'Capture method',.7,2.7,5,.4,22,INK,True)
text(s,'Flutter release web build; 393 x 852 mobile viewport; local seeded Laravel backend; supplied admin account. Original icon assets, with Arial fallback for network-loaded text fonts.',.7,3.28,5.45,2.1,18,GRAY)
text(s,'Presentation boundaries',7.0,2.7,5.5,.4,22,INK,True)
text(s,'Demo data is shown. Some flows use sample IDs or static states. External maps and services were not verified. Production deployment, clinical outcomes, security certification, traction, pricing, and integrations are not claimed.',7.0,3.28,5.45,2.5,18,GRAY)
text(s,'Included: route-indexed screenshot library, browser gallery, and editable PowerPoint deck.',.7,6.15,11.5,.55,18,GREEN,True)

prs.save(OUT)
cards=''.join(f'<a href="screenshots/{html.escape(r["file"])}" target="_blank"><img loading="lazy" src="screenshots/{html.escape(r["file"])}"><b>{html.escape(r["label"])}</b><small>{html.escape(r["route"])}</small></a>' for r in rows if 'file' in r)
(ROOT/'Screenshot_Gallery.html').write_text('''<!doctype html><html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SchooKeep Screenshot Library</title><style>body{margin:0;padding:28px;background:#eff5f2;color:#142822;font-family:Arial}h1{font-size:32px}p{max-width:850px;line-height:1.5}main{display:grid;grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:24px}a{display:block;text-decoration:none;color:inherit;background:white;padding:12px;border:1px solid #d5e2dc;border-radius:6px}img{width:100%;height:420px;object-fit:contain;background:#f7f9fb}b,small{display:block;margin-top:10px;overflow-wrap:anywhere}small{color:#5b7069}</style><h1>SchooKeep Screenshot Library</h1><p>145 routes captured from the Flutter app with local seeded data. Admin access uses the supplied demo account. Text fonts use Arial as a local fallback. Dynamic route samples may show empty or unavailable states. Select a screen to open the full-resolution image.</p><main>'''+cards+'</main></html>',encoding='utf8')
gallery=ROOT/'Screenshot_Gallery.html'
gallery.write_text(gallery.read_text(encoding='utf8').replace('145 routes',f'{capture_count} routes'),encoding='utf8')
with zipfile.ZipFile(ROOT/'SchooKeep_Screenshots.zip','w',zipfile.ZIP_DEFLATED) as z:
    for f in SHOTS.glob('*.png'):z.write(f,'screenshots/'+f.name)
    z.write(ROOT/'Screenshot_Gallery.html','Screenshot_Gallery.html')
    z.write(ROOT/'screenshots.json','screenshots.json')
print(f'Created {len(prs.slides)} slides and {len(rows)} screenshot entries: {OUT}')
